import json
from collections import defaultdict
from pathlib import Path

import openpyxl

path = Path(__file__).with_name("Test Case.xlsx")
wb = openpyxl.load_workbook(path, data_only=True)

hier = []
for row in wb["Hierarchy"].iter_rows(min_row=2, values_only=True):
    if row[3]:
        hier.append(
            {"Region": row[0], "Country": row[1], "Sector": row[2], "AssetID": row[3]}
        )

bs = defaultdict(lambda: {"YTD": 0.0, "Fcst": 0.0, "OP": 0.0, "sort": None})
for cat, aid, ytd, fcst, op, srt in wb["BS by Region"].iter_rows(
    min_row=2, values_only=True
):
    if not cat:
        continue
    key = (cat, aid)
    bs[key]["YTD"] += float(ytd or 0)
    bs[key]["Fcst"] += float(fcst or 0)
    bs[key]["OP"] += float(op or 0)
    bs[key]["sort"] = srt

cat_tot = defaultdict(lambda: {"YTD": 0.0, "Fcst": 0.0, "OP": 0.0})
for (cat, aid), v in bs.items():
    if cat == "AUM":
        continue
    for k in ("YTD", "Fcst", "OP"):
        cat_tot[cat][k] += v[k]

aum = {
    k: cat_tot["Equity"][k] + cat_tot["Debt"][k] + cat_tot["3rd Party"][k]
    for k in ("YTD", "Fcst", "OP")
}

aid_sec = {h["AssetID"]: h["Sector"] for h in hier}
aid_cty = {h["AssetID"]: h["Country"] for h in hier}

gav_sector = defaultdict(lambda: {"YTD": 0.0, "Fcst": 0.0, "OP": 0.0})
aum_country = defaultdict(lambda: {"YTD": 0.0, "Fcst": 0.0, "OP": 0.0})
equity_asset = defaultdict(lambda: {"YTD": 0.0, "Fcst": 0.0, "OP": 0.0})
for (cat, aid), v in bs.items():
    if cat == "Real Estate":
        for k in ("YTD", "Fcst", "OP"):
            gav_sector[aid_sec[aid]][k] += v[k]
    if cat in ("Equity", "Debt", "3rd Party"):
        for k in ("YTD", "Fcst", "OP"):
            aum_country[aid_cty[aid]][k] += v[k]
    if cat == "Equity":
        for k in ("YTD", "Fcst", "OP"):
            equity_asset[aid][k] += v[k]

ws = wb["KPI's"]
headers = [c.value for c in next(ws.iter_rows(min_row=17, max_row=17))]
kpi = defaultdict(lambda: defaultdict(float))
for row in ws.iter_rows(min_row=18, values_only=True):
    aid = row[0]
    if not aid:
        continue
    for i, h in enumerate(headers):
        if i < 2 or h is None or "Variance" in str(h):
            continue
        val = row[i]
        if val is None:
            continue
        kpi[aid][str(h)] += float(val)

kpi_tot = defaultdict(float)
for aid, cols in kpi.items():
    for h, val in cols.items():
        kpi_tot[h] += val

bridge = defaultdict(lambda: {"Bridge": 0.0, "Impact": 0.0, "sort": None})
for region, scenario, cat, br, impact, sort, _ in wb["Equity Bridge"].iter_rows(
    min_row=2, values_only=True
):
    key = (scenario, cat)
    bridge[key]["Bridge"] += float(br or 0)
    bridge[key]["Impact"] += float(impact or 0)
    bridge[key]["sort"] = sort

out = {
    "hier": hier,
    "cat_tot": dict(cat_tot),
    "aum": aum,
    "gav_sector": {k: dict(v) for k, v in gav_sector.items()},
    "aum_country": {k: dict(v) for k, v in aum_country.items()},
    "equity_asset": {k: dict(v) for k, v in equity_asset.items()},
    "kpi_tot": dict(kpi_tot),
    "kpi": {k: dict(v) for k, v in kpi.items()},
    "bridge_dec": {
        cat: v["Bridge"]
        for (sc, cat), v in sorted(bridge.items(), key=lambda x: x[1]["sort"] or 0)
        if sc == "Dec 2019"
    },
}

print(json.dumps(out, indent=2))
